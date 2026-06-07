from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_FILE = Path("praca_w_zespole_medycznym.pptx")
TITLE_COLOR = RGBColor(30, 60, 120)
TEXT_COLOR = RGBColor(40, 40, 40)


SLIDES = [
    {
        "type": "title",
        "title": "Praca w zespole medycznym",
        "subtitle": "Konflikty, współpraca i komunikacja",
        "footer": "Prezentacja ok. 25 minut",
    },
    {
        "title": "Cel prezentacji",
        "bullets": [
            ("Pokazać, dlaczego komunikacja i współpraca są ważne w medycynie.", 0),
            ("Omówić konflikty w zespole medycznym.", 0),
            ("Przeanalizować 3 artykuły naukowe.", 0),
            ("Pokazać rolę elektroradiologa.", 0),
            ("Przedstawić 3 krótkie przypadki z praktyki.", 0),
        ],
    },
    {
        "title": "Zespół medyczny — kto go tworzy?",
        "bullets": [
            ("Lekarze.", 0),
            ("Pielęgniarki i położne.", 0),
            ("Ratownicy medyczni.", 0),
            ("Elektroradiolodzy.", 0),
            ("Diagności laboratoryjni.", 0),
            ("Fizjoterapeuci.", 0),
            ("Rejestratorki medyczne.", 0),
            ("Cel wspólny: bezpieczeństwo pacjenta.", 0),
        ],
    },
    {
        "title": "Dlaczego temat jest ważny?",
        "bullets": [
            ("Medycyna to praca zespołowa.", 0),
            ("Jedna zła informacja może spowodować błąd.", 0),
            ("Konflikt spowalnia pracę.", 0),
            ("Brak współpracy zwiększa stres.", 0),
            ("Dobra komunikacja poprawia bezpieczeństwo pacjenta.", 0),
        ],
    },
    {"type": "section", "title": "Analiza artykułów naukowych"},
    {
        "title": "Artykuł 1 — konflikty w zespole",
        "bullets": [
            ("Almost J. i wsp.", 0),
            ("Managing and mitigating conflict in healthcare teams: an integrative review", 0),
            ("DOI: 10.1111/jan.12903", 0),
            ("Temat: konflikty w zespołach medycznych.", 0),
            ("Rodzaj pracy: integracyjny przegląd literatury.", 0),
        ],
    },
    {
        "title": "Artykuł 1 — cel badania",
        "bullets": [
            ("Problem badawczy:", 0),
            ("Skąd biorą się konflikty w zespołach medycznych?", 1),
            ("Jak można je ograniczać?", 1),
            ("Dlaczego ważne?", 0),
            ("Konflikt pogarsza atmosferę pracy.", 1),
            ("Może utrudniać opiekę nad pacjentem.", 1),
            ("Może zwiększać ryzyko błędów.", 1),
        ],
    },
    {
        "title": "Artykuł 1 — metoda i grupa",
        "bullets": [
            ("Metoda: analiza wcześniejszych badań.", 0),
            ("Nie badano jednej konkretnej grupy pacjentów.", 0),
            ("Analizowano publikacje o konfliktach w ochronie zdrowia.", 0),
            ("Zakres: badania z lat 2002–2014.", 0),
            ("Liczba analizowanych artykułów: 44.", 0),
        ],
    },
    {
        "title": "Artykuł 1 — wyniki",
        "bullets": [
            ("Najczęstsze przyczyny konfliktów:", 0),
            ("słaba komunikacja,", 1),
            ("niejasne role,", 1),
            ("brak wsparcia,", 1),
            ("stres i presja czasu,", 1),
            ("trudne warunki pracy,", 1),
            ("brak szacunku.", 1),
        ],
    },
    {
        "title": "Artykuł 1 — znaczenie i ograniczenia",
        "bullets": [
            ("Znaczenie dla praktyki:", 0),
            ("trzeba jasno dzielić obowiązki,", 1),
            ("trzeba rozmawiać spokojnie,", 1),
            ("trzeba reagować wcześnie na konflikt.", 1),
            ("Ograniczenia:", 0),
            ("to przegląd literatury, nie jedno nowe badanie,", 1),
            ("mało danych o skutecznych interwencjach.", 1),
        ],
    },
    {
        "title": "Artykuł 2 — współpraca interdyscyplinarna",
        "bullets": [
            ("Schot E., Tummers L., Noordegraaf M.", 0),
            ("Working on working together.", 0),
            ("DOI: 10.1080/13561820.2019.1636007", 0),
            ("Temat: jak profesjonaliści medyczni tworzą współpracę.", 0),
            ("Rodzaj pracy: systematyczny przegląd literatury.", 0),
        ],
    },
    {
        "title": "Artykuł 2 — cel badania",
        "bullets": [
            ("Problem badawczy:", 0),
            ("Jak pracownicy ochrony zdrowia budują współpracę między zawodami?", 1),
            ("Dlaczego ważne?", 0),
            ("Pacjent często wymaga wielu specjalistów.", 1),
            ("Bez współpracy powstaje chaos.", 1),
            ("Dobra współpraca przyspiesza diagnostykę i leczenie.", 1),
        ],
    },
    {
        "title": "Artykuł 2 — metoda i grupa",
        "bullets": [
            ("Metoda: systematyczny przegląd literatury.", 0),
            ("Analizowano badania o współpracy między zawodami medycznymi.", 0),
            ("Bazy: Scopus, Web of Science, Medline.", 0),
            ("Liczba badań: 64.", 0),
            ("Wiele badań dotyczyło szpitali i zespołów klinicznych.", 0),
        ],
    },
    {
        "title": "Artykuł 2 — wyniki",
        "bullets": [
            ("Autorzy opisali 3 główne działania:", 0),
            ("1. Bridging gaps — łączenie luk informacyjnych.", 0),
            ("2. Negotiating overlaps — ustalanie granic ról.", 0),
            ("3. Creating spaces — tworzenie miejsca do rozmowy.", 0),
            ("Wniosek: współpraca nie dzieje się sama.", 0),
            ("Współpracę trzeba codziennie budować.", 0),
        ],
    },
    {
        "title": "Artykuł 2 — znaczenie i ograniczenia",
        "bullets": [
            ("Znaczenie dla praktyki:", 0),
            ("krótkie odprawy,", 1),
            ("jasna komunikacja,", 1),
            ("pytanie, gdy coś jest niejasne,", 1),
            ("szacunek do innych zawodów.", 1),
            ("Ograniczenia:", 0),
            ("dużo badań jakościowych,", 1),
            ("część wyników trudno uogólnić na wszystkie szpitale.", 1),
        ],
    },
    {
        "title": "Artykuł 3 — komunikacja i przekazywanie pacjenta",
        "bullets": [
            ("Arora V. i wsp.", 0),
            ("Communication failures in patient sign-out and suggestions for improvement.", 0),
            ("DOI: 10.1136/qshc.2005.015107", 0),
            ("Temat: błędy komunikacji przy przekazywaniu pacjenta.", 0),
            ("Rodzaj pracy: analiza incydentów krytycznych.", 0),
        ],
    },
    {
        "title": "Artykuł 3 — cel badania",
        "bullets": [
            ("Problem badawczy:", 0),
            ("Jak błędy w przekazywaniu informacji wpływają na bezpieczeństwo pacjenta?", 1),
            ("Dlaczego ważne?", 0),
            ("Pominięta informacja może spowodować złą decyzję.", 1),
            ("Może dojść do opóźnienia leczenia.", 1),
            ("Może dojść do niepotrzebnych badań lub błędów.", 1),
        ],
    },
    {
        "title": "Artykuł 3 — metoda i grupa",
        "bullets": [
            ("Metoda: critical incident technique.", 0),
            ("Badano konkretne trudne sytuacje.", 0),
            ("Grupa: 26 lekarzy stażystów/rezydentów pierwszego roku.", 0),
            ("Dotyczyło przekazywania informacji o pacjentach.", 0),
            ("Analizowano sytuacje, w których komunikacja była niewystarczająca.", 0),
        ],
    },
    {
        "title": "Artykuł 3 — wyniki",
        "bullets": [
            ("Opisano liczne incydenty wynikające ze złej komunikacji.", 0),
            ("Najczęstsze problemy:", 0),
            ("brak informacji o lekach,", 1),
            ("brak informacji o aktywnych problemach,", 1),
            ("brak informacji o badaniach w toku,", 1),
            ("nieaktualne notatki,", 1),
            ("brak rozmowy twarzą w twarz.", 1),
        ],
    },
    {
        "title": "Artykuł 3 — znaczenie i ograniczenia",
        "bullets": [
            ("Znaczenie dla praktyki:", 0),
            ("informacje trzeba przekazywać jasno,", 1),
            ("dokumentacja musi być aktualna,", 1),
            ("nie wolno zgadywać.", 1),
            ("Ograniczenia:", 0),
            ("badanie dotyczyło konkretnej grupy lekarzy,", 1),
            ("wyniki nie muszą pasować do każdego oddziału.", 1),
        ],
    },
    {"type": "section", "title": "Opracowanie tematu"},
    {
        "title": "Najczęstsze konflikty w zespole medycznym",
        "bullets": [
            ("Lekarz — pielęgniarka.", 0),
            ("Lekarz — elektroradiolog.", 0),
            ("Personel — pacjent.", 0),
            ("Personel — rodzina pacjenta.", 0),
            ("Konflikt między zmianami dyżurowymi.", 0),
            ("Konflikt przez niejasne polecenia.", 0),
        ],
    },
    {
        "title": "Przyczyny konfliktów",
        "bullets": [
            ("Stres.", 0),
            ("Presja czasu.", 0),
            ("Zmęczenie.", 0),
            ("Brak personelu.", 0),
            ("Niejasny podział obowiązków.", 0),
            ("Niepełna dokumentacja.", 0),
            ("Brak szacunku.", 0),
            ("Zła komunikacja.", 0),
        ],
    },
    {
        "title": "Jak poprawić współpracę?",
        "bullets": [
            ("Mówić jasno i krótko.", 0),
            ("Potwierdzać ważne informacje.", 0),
            ("Pytać, gdy coś jest niejasne.", 0),
            ("Nie obwiniać od razu.", 0),
            ("Znać swoją rolę.", 0),
            ("Szanować innych członków zespołu.", 0),
            ("Stosować procedury.", 0),
            ("Robić krótkie odprawy.", 0),
        ],
    },
    {
        "title": "Rola elektroradiologa",
        "bullets": [
            ("Sprawdza dane pacjenta.", 0),
            ("Sprawdza skierowanie.", 0),
            ("Przygotowuje pacjenta do badania.", 0),
            ("Dba o bezpieczeństwo radiologiczne.", 0),
            ("Komunikuje się z lekarzem i pielęgniarką.", 0),
            ("Zgłasza niejasności.", 0),
            ("Nie wykonuje badania „na domysł”.", 0),
            ("Chroni pacjenta i siebie.", 0),
        ],
    },
    {"type": "section", "title": "Mini-przypadki"},
    {
        "title": "Mini-przypadek 1 — nieczytelne skierowanie",
        "bullets": [
            ("Sytuacja:", 0),
            ("Pacjent przychodzi na RTG. Skierowanie jest nieczytelne.", 1),
            ("Problem:", 0),
            ("Ryzyko wykonania złego badania.", 1),
            ("Co robi elektroradiolog?", 0),
            ("Nie zgaduje.", 1),
            ("Kontaktuje się z lekarzem.", 1),
            ("Wyjaśnia pacjentowi sytuację.", 1),
            ("Wykonuje badanie dopiero po potwierdzeniu.", 1),
        ],
    },
    {
        "title": "Mini-przypadek 2 — TK z kontrastem",
        "bullets": [
            ("Sytuacja:", 0),
            ("Lekarz chce pilnie TK z kontrastem.", 1),
            ("Brakuje aktualnej kreatyniny.", 1),
            ("Problem:", 0),
            ("Presja czasu kontra bezpieczeństwo pacjenta.", 1),
            ("Co robi elektroradiolog?", 0),
            ("Mówi spokojnie o braku danych.", 1),
            ("Proponuje szybkie uzupełnienie informacji.", 1),
            ("Kontaktuje lekarza radiologa lub przełożonego, jeśli trzeba.", 1),
        ],
    },
    {
        "title": "Mini-przypadek 3 — agresywny pacjent",
        "bullets": [
            ("Sytuacja:", 0),
            ("Pacjent krzyczy i odmawia współpracy.", 1),
            ("Problem:", 0),
            ("Badanie może być niebezpieczne.", 1),
            ("Co robi elektroradiolog?", 0),
            ("Mówi spokojnie.", 1),
            ("Krótko tłumaczy badanie.", 1),
            ("Prosi o pomoc pielęgniarkę lub lekarza.", 1),
            ("Nie wykonuje badania na siłę.", 1),
            ("Dba o bezpieczeństwo wszystkich.", 1),
        ],
    },
    {
        "title": "Porównanie artykułów",
        "bullets": [
            ("Almost i wsp.: konflikt wynika ze złej komunikacji, stresu i niejasnych ról.", 0),
            ("Schot i wsp.: współpracę trzeba aktywnie budować.", 0),
            ("Arora i wsp.: złe przekazanie informacji może zaszkodzić pacjentowi.", 0),
            ("Wspólny wniosek:", 0),
            ("mówić jasno,", 1),
            ("znać swoją rolę,", 1),
            ("reagować szybko,", 1),
            ("nie zgadywać.", 1),
        ],
    },
    {
        "title": "Wnioski końcowe",
        "bullets": [
            ("Dobra komunikacja zmniejsza liczbę błędów.", 0),
            ("Konflikty są częste, ale można je rozwiązywać.", 0),
            ("Współpraca poprawia bezpieczeństwo pacjenta.", 0),
            ("Elektroradiolog jest ważną częścią zespołu.", 0),
            ("Najważniejsze: nie milczeć, nie zgadywać, pytać.", 0),
        ],
    },
    {
        "title": "Bibliografia",
        "bullets": [
            ("Almost J. i wsp. Managing and mitigating conflict in healthcare teams: an integrative review. DOI: 10.1111/jan.12903", 0),
            ("Schot E., Tummers L., Noordegraaf M. Working on working together. DOI: 10.1080/13561820.2019.1636007", 0),
            ("Arora V. i wsp. Communication failures in patient sign-out and suggestions for improvement. DOI: 10.1136/qshc.2005.015107", 0),
            ("Linki DOI:", 0),
            ("https://doi.org/10.1111/jan.12903", 0),
            ("https://doi.org/10.1080/13561820.2019.1636007", 0),
            ("https://doi.org/10.1136/qshc.2005.015107", 0),
        ],
    },
]


def style_run(run, font_size, bold=False):
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = TITLE_COLOR if bold else TEXT_COLOR


def add_textbox(slide, left, top, width, height):
    textbox = slide.shapes.add_textbox(left, top, width, height)
    frame = textbox.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return frame


def add_title_slide(prs, title, subtitle, footer):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_frame = add_textbox(slide, Inches(0.8), Inches(1.2), Inches(11.8), Inches(1.3))
    title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = title_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = title
    style_run(run, 24, bold=True)

    subtitle_frame = add_textbox(slide, Inches(1.2), Inches(2.7), Inches(11.0), Inches(0.8))
    subtitle_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = subtitle_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = subtitle
    style_run(run, 20)

    footer_frame = add_textbox(slide, Inches(3.7), Inches(6.5), Inches(6.0), Inches(0.4))
    paragraph = footer_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = footer
    style_run(run, 12)


def add_section_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title_frame = add_textbox(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.2))
    title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = title_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = title
    style_run(run, 26, bold=True)


def add_content_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_frame = add_textbox(slide, Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.65))
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.alignment = PP_ALIGN.LEFT
    title_run = title_paragraph.add_run()
    title_run.text = title
    style_run(title_run, 20, bold=True)

    body_frame = add_textbox(slide, Inches(0.9), Inches(1.2), Inches(11.6), Inches(5.8))
    for index, (text, level) in enumerate(bullets):
        paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
        paragraph.level = level
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(4)
        paragraph.line_spacing = 1.08
        run = paragraph.add_run()
        run.text = text
        style_run(run, 15)


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for slide_data in SLIDES:
        slide_type = slide_data.get("type", "content")
        if slide_type == "title":
            add_title_slide(prs, slide_data["title"], slide_data["subtitle"], slide_data["footer"])
        elif slide_type == "section":
            add_section_slide(prs, slide_data["title"])
        else:
            add_content_slide(prs, slide_data["title"], slide_data["bullets"])

    prs.save(OUTPUT_FILE)


if __name__ == "__main__":
    build_presentation()
